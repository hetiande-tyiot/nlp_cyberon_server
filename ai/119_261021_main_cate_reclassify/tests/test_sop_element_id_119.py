"""已填欄位回傳 SOP 要素編號（次案類編號-題號）對抽到的內容，例如 {"202-1": "機車追撞"}。"""

from __future__ import annotations

import unittest

from case_info_119 import CaseInfo119


class SopElementIdTests(unittest.TestCase):
    def test_traffic_accident_injury_cause_is_202_1(self) -> None:
        case = CaseInfo119(
            main_category="救護",
            sub_category="車禍",
            injury_cause="機車追撞",
            incident_description="被車撞到",
            address="板橋區文化路一段1號",
        )

        ids = case.to_dict()["sop_element_ids"]

        # 只有受傷原因對得上車禍的題目；發生什麼事、地址不在車禍的關鍵要素裡
        self.assertEqual(ids, {"202-1": "機車追撞"})

    def test_same_label_uses_the_subtype_own_number(self) -> None:
        case = CaseInfo119(
            main_category="救護",
            sub_category="一般受傷",
            injury_cause="跌倒",
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["210-1"], "跌倒")
        self.assertNotIn("202-1", ids)

    def test_fire_townhouse_smoke_is_101_5(self) -> None:
        case = CaseInfo119(
            main_category="火警",
            sub_category="透天厝",
            fire_or_smoke="有火有煙",
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["101-5"], "有火有煙")  # 1013 版透天厝第 5 題：火煙狀況

    def test_yes_no_answers_are_written_as_text(self) -> None:
        case = CaseInfo119(
            main_category="救護",
            sub_category="一般受傷",
            consciousness=True,
            caller_is_patient=False,
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["210-8"], "是")  # 一般受傷第 8 題：意識
        self.assertEqual(ids["210-3"], "否")  # 一般受傷第 3 題：報案人與患者關係

    def test_building_structure_code_is_skipped_for_text_version(self) -> None:
        # building_structure 是派遣卡數字代碼（2＝鐵皮），同一題用文字的 building_construction
        case = CaseInfo119(
            main_category="火警",
            sub_category="透天厝",
            building_structure=2,
            building_construction="鐵皮屋",
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["101-13"], "鐵皮屋")  # 1013 版透天厝第 13 題：建物構造

    def test_unmatched_field_and_summary_omit_the_map(self) -> None:
        case = CaseInfo119(
            main_category="救護",
            sub_category="車禍",
            injury_cause="機車追撞",
            current_condition="現場安全",
        )

        ids = case.to_dict()["sop_element_ids"]
        summary_fields = case.filled_fields_for_summary()

        # 目前狀況不是車禍的關鍵要素，不給編號
        self.assertEqual(list(ids), ["202-1"])
        self.assertNotIn("sop_element_ids", summary_fields)
        self.assertEqual(summary_fields["injury_cause"], "機車追撞")

    def test_fire_lists_questions_still_to_ask_as_pending(self) -> None:
        # 住宅、看得到火：非工廠的建築物類型、濃煙顏色、起火樓層要問；工廠才問的不列
        case = CaseInfo119(
            main_category="火警",
            sub_category="透天厝",
            place_usage="住宅",
            fire_or_smoke="有火有煙",
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["101-1"], "住宅")
        self.assertEqual(ids["101-2"], "待問")  # 建築物類型（非工廠）：要問，還沒答
        self.assertNotIn("101-3", ids)  # 建築物類型（工廠）：不用問
        self.assertEqual(ids["101-6"], "待問")  # 濃煙顏色：看得到煙才問
        self.assertEqual(ids["101-7"], "待問")  # 起火樓層：非工廠、看得到火才問
        self.assertNotIn("101-8", ids)  # 燃燒物：工廠才問
        self.assertNotIn("101-15", ids)  # 有無爆炸：被動題，沒人講到就不列

    def test_fire_factory_answer_goes_to_factory_question(self) -> None:
        # 同一個欄位有兩題（非工廠、工廠）時，答案放在前置條件成立的那一題
        case = CaseInfo119(
            main_category="火警",
            sub_category="透天厝",
            place_usage="鐵皮工廠",
            building_type="大型違章建築區與傳統市場(含工廠)",
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["101-3"], "大型違章建築區與傳統市場(含工廠)")
        self.assertNotIn("101-2", ids)

    def test_fire_conditions_not_met_are_left_out(self) -> None:
        # 垃圾、無火無煙：只有「氣味」要問；濃煙顏色、是否延燒都不列
        case = CaseInfo119(
            main_category="火警",
            sub_category="垃圾",
            fire_or_smoke="無火無煙",
        )

        ids = case.to_dict()["sop_element_ids"]

        self.assertEqual(ids["130-9"], "待問")  # 氣味
        self.assertNotIn("130-7", ids)  # 濃煙顏色
        self.assertNotIn("130-10", ids)  # 是否延燒

    def test_fire_tab_questions_match_case_type_qa_order(self) -> None:
        # 火警的編號用「垂片第幾題」當題號，所以垂片題目要跟 CaseTypeQA.xlsx 一題一題對得上。
        # 對不上代表改了其中一邊沒同步另一邊。
        from fire_tab_map_119 import TAB_QUESTIONS
        from sop_element_id_119 import _XLSX, _SHEET
        from openpyxl import load_workbook

        # 10/08 統一的關鍵要素名稱：垂片照 xlsx 原文，CaseTypeQA 用統一後的名稱
        unified = {
            "延燒可能": "是否延燒",
            "起火戶應門": "應門或聯絡狀況",
            "無人應門或聯絡不上": "應門或聯絡狀況",
            "氣味類型": "氣味",
            "延燒面積": "燃燒範圍",
            "燃燒面積": "燃燒範圍",
        }
        rows: dict = {}
        sheet = load_workbook(_XLSX, read_only=True)[_SHEET]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            rows.setdefault(row[1], []).append(row[5])  # 次案類 → 關鍵要素（依題號）
        for tab, subtype in (("A", "火警-透天厝"), ("B1", "火警-汽車"),
                             ("B2", "火警-山林田野(平地)"), ("C", "火警-垃圾")):
            with self.subTest(tab=tab):
                expected = [unified.get(q.element, q.element) for q in TAB_QUESTIONS[tab]]
                self.assertEqual(rows[subtype], expected)

    def test_empty_case_has_no_ids(self) -> None:
        self.assertEqual(CaseInfo119().to_dict()["sop_element_ids"], {})


if __name__ == "__main__":
    unittest.main()
