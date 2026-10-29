"""fire_level_119.py — 火警等級：次案類初始等級，語意標籤只升不降。

初始等級來自 scripts/levels.xlsx「火災分級-修正」第 2–31 行（C 次分類、N 初始等級）。
空白、資料庫、X 都是 0。升級只看 16 個固定語意，不從其他火警欄位換算。
算出來的等級寫進案件的 level（派遣等級，各大類共用的欄位；其他大類目前固定 0）。
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence

_XLSX = Path(__file__).resolve().parent / "scripts" / "levels.xlsx"
_SHEET = "火災分級-修正"
_FIRST_ROW = 2
_LAST_ROW = 31

# 代碼裡的次案類名稱 → 表上 C 列的寫法
_SUBCATEGORY_ALIASES = {
    "運輸中樞用途建築物": "運輸中樞建築物",
    "學校、醫院、老人院建築物": "醫院老人院建築物",
    "汽車": "汽車火警",
    "機車": "機車火警",
    "隧道": "隧道火警",
    "軌道型交通工具": "軌道型交通工具火警",
    "船舶": "船帕",
}

_LEVEL_TEXT = {"一級": 1, "二級": 2, "三級": 3}
_IGNORED_INITIAL = {"", "資料庫", "X"}

# 語意標籤 → 要升到的等級。同一組互相排斥，prompt 裡要求只留一個。
SIGNAL_LEVEL: Dict[str, int] = {
    "有火焰": 2,
    "無火焰": 1,
    "黑煙": 2,
    "白煙": 1,
    "其他色煙": 1,
    "無煙": 1,
    "有人受困": 3,
    "無人受困": 1,
    "極可能或已延燒": 3,
    "延燒可能性低": 2,
    "未知面積": 1,
    "0~50坪/165平方公尺": 1,
    "50~100坪/330平方公尺": 2,
    "100~300坪/1000平方公尺": 2,
    "300~500坪/1500平方公尺": 3,
    "500坪以上": 3,
}

SIGNAL_LABELS: tuple[str, ...] = tuple(SIGNAL_LEVEL)

_EXCLUSIVE_GROUPS: tuple[tuple[str, ...], ...] = (
    ("有火焰", "無火焰"),
    ("黑煙", "白煙", "其他色煙", "無煙"),
    ("有人受困", "無人受困"),
    ("極可能或已延燒", "延燒可能性低"),
    (
        "未知面積",
        "0~50坪/165平方公尺",
        "50~100坪/330平方公尺",
        "100~300坪/1000平方公尺",
        "300~500坪/1500平方公尺",
        "500坪以上",
    ),
)


def _parse_initial(value: Any) -> int:
    if value is None:
        return 0
    text = str(value).strip()
    if text in _IGNORED_INITIAL:
        return 0
    return _LEVEL_TEXT.get(text, 0)


@lru_cache(maxsize=1)
def load_initial_levels(path: str = str(_XLSX)) -> Dict[str, int]:
    """次分類名稱 → 初始等級 0–3。只讀第 2–31 行。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("讀取 levels.xlsx 需要安裝 openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if _SHEET not in workbook.sheetnames:
            raise ValueError(f"levels.xlsx 缺少工作表：{_SHEET}")
        sheet = workbook[_SHEET]
        rows = sheet.iter_rows(values_only=True)
        next(rows, None)
        table: Dict[str, int] = {}
        for excel_row, row in enumerate(rows, start=_FIRST_ROW):
            if excel_row > _LAST_ROW:
                break
            if not row or len(row) < 3:
                continue
            name = str(row[2] or "").strip()
            if not name:
                continue
            initial = _parse_initial(row[13] if len(row) > 13 else None)
            table[name] = initial
        return table
    finally:
        workbook.close()


def initial_level(sub_category: Optional[str]) -> int:
    """次案類的初始等級。對不上、或表上是空白／資料庫／X，都是 0。"""
    name = (sub_category or "").strip()
    if not name:
        return 0
    name = _SUBCATEGORY_ALIASES.get(name, name)
    return load_initial_levels().get(name, 0)


def accepted_signals(raw: Any) -> List[str]:
    """留下清單上的標籤。同一組互相排斥時只留等級較高的一個。"""
    if raw is None:
        return []
    if isinstance(raw, str):
        items: Iterable[Any] = [raw]
    elif isinstance(raw, Sequence) and not isinstance(raw, (bytes, bytearray)):
        items = raw
    else:
        return []
    found = []
    seen = set()
    for item in items:
        text = str(item or "").strip()
        if text in SIGNAL_LEVEL and text not in seen:
            seen.add(text)
            found.append(text)
    chosen: List[str] = []
    consumed = set()
    for group in _EXCLUSIVE_GROUPS:
        members = [label for label in found if label in group]
        if not members:
            continue
        best = max(members, key=lambda label: (SIGNAL_LEVEL[label], -group.index(label)))
        chosen.append(best)
        consumed.update(group)
    for label in found:
        if label not in consumed:
            chosen.append(label)
    return chosen


def highest_signal_level(raw: Any) -> int:
    signals = accepted_signals(raw)
    if not signals:
        return 0
    return max(SIGNAL_LEVEL[label] for label in signals)


def raise_fire_level(case: Any, signals: Any = None) -> None:
    """火警案件的 level 取 max(目前, 次案類初始, 本輪語意)。非火警不寫，維持 0。"""
    if getattr(case, "main_category", None) != "火警":
        return
    current = getattr(case, "level", None)
    if isinstance(current, bool) or not isinstance(current, int):
        current = 0
    updated = max(current, initial_level(getattr(case, "sub_category", None)), highest_signal_level(signals))
    case.level = updated
