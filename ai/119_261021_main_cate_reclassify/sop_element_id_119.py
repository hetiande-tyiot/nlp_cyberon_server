"""sop_element_id_119.py — 已填欄位 → SOP 要素編號與值。

編號格式是「次案類編號-題號」，例如救護-車禍的受傷原因是 ``202-1``。
回傳 ``{欄位: {"id": "202-1", "value": <已抽取的值>}}``。
對照表來自 scripts/CaseTypeQA.xlsx：用「主類別-次案類」加欄位中文名
對上「次案類 + 關鍵要素」。對不上的欄位不給編號。
"""

from __future__ import annotations

from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Mapping, Optional, Tuple

from case_field_labels_119 import LABELS

_XLSX = Path(__file__).resolve().parent / "scripts" / "CaseTypeQA.xlsx"
_SHEET = "次案類問題"

# 代碼裡的次案類名稱 → Excel「次案類」橫線後面的名稱
_SUBCATEGORY_ALIASES = {
    "孕婦急產": "孕婦待產",
}

# 欄位中文名（LABELS）→ Excel「關鍵要素」。用字相同的不必寫在這裡。
_LABEL_TO_ELEMENT = {
    "是否有意識": "意識",
    "是否有呼吸": "呼吸",
    "傷病患人數": "患者人數",
    "過去病史": "病史",
    "發生了什麼事": "病徵",
    "延燒可能": "是否延燒",
    "應門狀況（起火戶或負責人）": "應門或聯絡狀況",
    "報案人是否即患者": "報案人與患者關係",
    "宮縮情況": "陣痛狀況",
    "預產期": "預產期與懷孕週數",
    "產檢異常史": "產檢狀況",
    # LABELS 寫「看到火還是煙」，表裡的關鍵要素是「火煙狀況」（透天厝第 4 題）。
    "看到火還是煙": "火煙狀況",
}


def case_type_key(main_category: Optional[str], sub_category: Optional[str]) -> Optional[str]:
    """拼出 Excel「次案類」欄的鍵。次案類未定、主類已是救護時用表裡的「救護」。"""
    main = (main_category or "").strip()
    sub = (sub_category or "").strip()
    if sub:
        sub = _SUBCATEGORY_ALIASES.get(sub, sub)
        if not main:
            return None
        return f"{main}-{sub}"
    if main == "救護":
        return "救護"
    return None


def _is_filled(field: str, value: Any) -> bool:
    """與 CaseInfo119.is_field_filled 同一套空值判斷。False 布林算已填。"""
    if value is None:
        return False
    if field == "ImportantCase" and value == 0:
        return False
    if isinstance(value, str):
        return bool(value.strip())
    if isinstance(value, (list, dict)):
        return bool(value)
    return True


def _number_text(value: Any) -> str:
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value or "").strip()
    try:
        number = float(text)
    except ValueError:
        return text
    if number.is_integer():
        return str(int(number))
    return text


@lru_cache(maxsize=1)
def load_element_index(path: str = str(_XLSX)) -> Dict[Tuple[str, str], str]:
    """(次案類, 關鍵要素) → 「編號-題號」。同一組重複時保留題號較小的那列。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("讀取 CaseTypeQA.xlsx 需要安裝 openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if _SHEET not in workbook.sheetnames:
            raise ValueError(f"CaseTypeQA.xlsx 缺少工作表：{_SHEET}")
        sheet = workbook[_SHEET]
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell or "").strip() for cell in next(rows, ())]
        required = ("次案類編號", "次案類", "題號", "關鍵要素")
        missing = [name for name in required if name not in header]
        if missing:
            raise ValueError(f"CaseTypeQA.xlsx 缺少欄位：{', '.join(missing)}")
        index = {name: header.index(name) for name in required}
        table: Dict[Tuple[str, str], str] = {}
        for row in rows:
            if not row or row[index["次案類編號"]] is None:
                continue
            subtype = str(row[index["次案類"]] or "").strip()
            element = str(row[index["關鍵要素"]] or "").strip()
            if not subtype or not element:
                continue
            code = (
                f"{_number_text(row[index['次案類編號']])}"
                f"-{_number_text(row[index['題號']])}"
            )
            key = (subtype, element)
            current = table.get(key)
            if current is None or _question_number(code) < _question_number(current):
                table[key] = code
        return table
    finally:
        workbook.close()


def _question_number(code: str) -> int:
    try:
        return int(code.rsplit("-", 1)[1])
    except (IndexError, ValueError):
        return 10**9


def sop_element_ids(case: Mapping[str, Any]) -> Dict[str, Dict[str, Any]]:
    """已填且對得上表的欄位 → {id, value}。鍵順序與 LABELS 一致。"""
    subtype = case_type_key(case.get("main_category"), case.get("sub_category"))
    if not subtype:
        return {}
    table = load_element_index()
    found: Dict[str, Dict[str, Any]] = {}
    for field, label in LABELS.items():
        if field not in case or not _is_filled(field, case.get(field)):
            continue
        element = _LABEL_TO_ELEMENT.get(label, label)
        code = table.get((subtype, element))
        if code:
            found[field] = {"id": code, "value": case.get(field)}
    return found
