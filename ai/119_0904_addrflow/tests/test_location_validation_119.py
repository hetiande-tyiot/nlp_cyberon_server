from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

from location_validation_119 import (
    DEFAULT_LANDMARK_XLSX,
    JurisdictionResult,
    load_landmark_names,
    load_mrt_location_names,
    query_jurisdiction,
    validate_landmark,
    validate_mrt,
)


class LocationValidationTests(unittest.TestCase):
    def tearDown(self) -> None:
        load_landmark_names.cache_clear()
        load_mrt_location_names.cache_clear()

    def test_landmark_template_has_expected_columns(self) -> None:
        workbook = load_workbook(DEFAULT_LANDMARK_XLSX, read_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows())]
        workbook.close()
        self.assertEqual(headers, ["地標名稱", "別名", "地址"])

    def test_landmark_loader_supports_aliases(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "landmarks.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["地標名稱", "別名", "地址"])
            sheet.append(["板橋大觀市場", "大觀市場、板橋市場", "新北市板橋區"])
            workbook.save(path)

            self.assertIn("板橋市場", load_landmark_names(str(path)))
            self.assertTrue(validate_landmark("我在大觀市場旁邊", path=str(path)))
            self.assertFalse(validate_landmark("未知市場旁邊", path=str(path)))

    def test_real_mrt_csv_uses_second_column_cp950(self) -> None:
        names = load_mrt_location_names()
        self.assertIn("頂埔站出口1", names)
        self.assertTrue(validate_mrt("捷運頂埔站出口1"))
        self.assertFalse(validate_mrt("捷運不存在站出口9"))

    @patch("location_validation_119.urlopen")
    def test_jurisdiction_requires_nonempty_office_name(
        self, mocked_urlopen: MagicMock
    ) -> None:
        response = MagicMock()
        response.read.return_value = (
            b'{"status":"Success","officeName":"Xindian Office"}'
        )
        mocked_urlopen.return_value.__enter__.return_value = response

        result = query_jurisdiction("新北市新店區中央路133巷")

        self.assertEqual(
            result,
            JurisdictionResult(True, office_name="Xindian Office"),
        )
        request = mocked_urlopen.call_args[0][0]
        self.assertIn(
            "http://100.107.145.7:8088/api/AddrCheck/Verify",
            request.full_url,
        )

    @patch("location_validation_119.urlopen")
    def test_jurisdiction_empty_office_is_invalid(
        self, mocked_urlopen: MagicMock
    ) -> None:
        response = MagicMock()
        response.read.return_value = b'{"status":"Success","officeName":null}'
        mocked_urlopen.return_value.__enter__.return_value = response
        self.assertIs(query_jurisdiction("不存在地址").valid, False)

    @patch("location_validation_119.urlopen", side_effect=TimeoutError("timeout"))
    def test_jurisdiction_error_is_distinct_from_invalid(
        self, _mocked_urlopen: MagicMock
    ) -> None:
        result = query_jurisdiction("新北市板橋區")
        self.assertIsNone(result.valid)
        self.assertIn("timeout", result.error or "")


if __name__ == "__main__":
    unittest.main()
