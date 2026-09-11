"""119 地址类别的数据加载与外部辖区校验。"""

from __future__ import annotations

import csv
import json
import os
import re
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Iterable, Optional
from urllib.parse import urlencode
from urllib.request import Request, urlopen


BASE_DIR = Path(__file__).resolve().parent
DEFAULT_JURISDICTION_API_URL = (
    "http://61.216.64.19:8055/Template/Office/data"
)
DEFAULT_MRT_CSV = BASE_DIR / "scripts" / "MRTstation.csv"
DEFAULT_LANDMARK_XLSX = BASE_DIR / "scripts" / "landmarks.xlsx"

# ── addrCheck 地址驗測（2026-09-01 起地址確認全面改走 addrCheck /Verify）──
# env：ADDRCHECK_API_URL / ADDRCHECK_API_TOKEN（Bearer，與 ACRC 共用 Ingest__Token）
#      / ADDRCHECK_API_TIMEOUT（預設 2.5）
DEFAULT_ADDRCHECK_BASE_URL = "http://192.168.5.120:8088"
ADDRCHECK_VERIFY_PATH = "/api/AddrCheck/Verify"

# 內部地點型態 → addrCheck API 的 type（帶 type 只查該層；None 讓後端自動判斷）
LOCATION_TYPE_TO_API_TYPE = {
    "address": "House",
    "intersection": "Crossroad",
    "highway": "Freeway",
    "landmark": "Landmark",
    "mrt": "Landmark",
}


@dataclass(frozen=True)
class JurisdictionResult:
    """辖区 API 结果；valid=None 表示网络或响应格式异常。"""

    valid: Optional[bool]
    office_name: Optional[str] = None
    error: Optional[str] = None
    api_reason: Optional[str] = None  # addrCheck 回傳之 reason（road_only/not_found…）


def _extract_office_name(payload: Any) -> Optional[str]:
    if isinstance(payload, dict):
        value = payload.get("officeName")
        if value is not None and str(value).strip():
            return str(value).strip()
        for key in ("data", "result"):
            nested = payload.get(key)
            found = _extract_office_name(nested)
            if found:
                return found
    elif isinstance(payload, list):
        for item in payload:
            found = _extract_office_name(item)
            if found:
                return found
    return None


# addrCheck 失敗分類 → 值班看板上的錯誤說明
ADDRCHECK_REASON_TEXT = {
    "road_only": "查無此門牌",
    "not_found": "查無此路段",
    "invalid_format": "地址格式無法解析",
}
DEFAULT_ADDRCHECK_FAILURE_TEXT = "地址驗測查無此地點"


def addrcheck_failure_reason(api_reason: Optional[str]) -> str:
    """把 addrCheck 的 reason 轉成可讀的錯誤說明。"""
    return ADDRCHECK_REASON_TEXT.get(
        (api_reason or "").strip(), DEFAULT_ADDRCHECK_FAILURE_TEXT
    )


def _addrcheck_token() -> str:
    """addrCheck Bearer token（與 ACRC 共用 Ingest__Token）。"""
    for key in ("ADDRCHECK_API_TOKEN", "INGEST_TOKEN", "Ingest__Token"):
        value = os.getenv(key)
        if value:
            return value
    return ""


def verify_address_status(
    address: str,
    location_type: Optional[str] = None,
    *,
    url: Optional[str] = None,
    timeout: Optional[float] = None,
) -> tuple[Optional[bool], Optional[str]]:
    """地址驗測：POST addrCheck /Verify，回 (status, hint)。

    status True=有效 / False=查無 / None=網路或回應異常（API 無法使用）。
    需要 API 的 reason 分類時改用 verify_address_detail。
    """
    status, hint, _reason = verify_address_detail(
        address, location_type, url=url, timeout=timeout
    )
    return status, hint


def verify_address_detail(
    address: str,
    location_type: Optional[str] = None,
    *,
    url: Optional[str] = None,
    timeout: Optional[float] = None,
) -> tuple[Optional[bool], Optional[str], Optional[str]]:
    """地址驗測核心：POST addrCheck /Verify，回 (status, hint, reason)。

    reason 為 API 的失敗分類（road_only=查無門牌 / not_found=查無路段…），
    供上層產生正確的錯誤說明與覆問話術。
    location_type 轉成 API 的 type 只查該層；None 讓後端自動判斷。
    """
    address = (address or "").strip()
    if not address:
        return False, "尚未取得地址", None
    base = (url or os.getenv("ADDRCHECK_API_URL", DEFAULT_ADDRCHECK_BASE_URL)).rstrip("/")
    endpoint = f"{base}{ADDRCHECK_VERIFY_PATH}"
    if timeout is None:
        try:
            timeout = float(os.getenv("ADDRCHECK_API_TIMEOUT", "2.5"))
        except ValueError:
            timeout = 2.5
    body: dict[str, object] = {"address": address}
    api_type = LOCATION_TYPE_TO_API_TYPE.get(location_type) if location_type else None
    if api_type:
        body["type"] = api_type
    headers = {"Content-Type": "application/json", "Accept": "application/json"}
    token = _addrcheck_token()
    if token:
        headers["Authorization"] = f"Bearer {token}"
    request = Request(
        endpoint,
        data=json.dumps(body, ensure_ascii=False).encode("utf-8"),
        headers=headers,
        method="POST",
    )
    try:
        with urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        return None, str(exc), None
    if not isinstance(payload, dict):
        return None, "回應格式非物件", None
    return (
        bool(payload.get("status")),
        payload.get("hint"),
        payload.get("reason"),
    )


def query_jurisdiction(
    address: str,
    *,
    url: Optional[str] = None,
    timeout: Optional[float] = None,
) -> JurisdictionResult:
    """一般地址（address 型態）驗測 → addrCheck House。

    回傳沿用 JurisdictionResult 維持引擎介面：valid=status，office_name 固定 None
    （119 不需管轄分局），valid=None 表示 API 無法使用。
    """
    status, hint, reason = verify_address_detail(
        address, "address", url=url, timeout=timeout
    )
    return JurisdictionResult(
        status, office_name=None, error=hint, api_reason=reason
    )


def _normalized_name(value: str) -> str:
    return re.sub(r"[\s　，,。．、;；/／()（）]+", "", str(value or "")).lower()


def _split_aliases(value: str) -> Iterable[str]:
    for item in re.split(r"[、,，;；/／]+", str(value or "")):
        item = item.strip()
        if item:
            yield item


@lru_cache(maxsize=4)
def load_landmark_names(path: str = str(DEFAULT_LANDMARK_XLSX)) -> tuple[str, ...]:
    """读取地标名称与别名；模板为空时返回空集合。"""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("讀取 landmarks.xlsx 需要安裝 openpyxl") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(v or "").strip() for v in next(rows, ())]
    required = ["地標名稱", "別名", "地址"]
    if headers[:3] != required:
        workbook.close()
        raise ValueError(f"地標 Excel 欄位必須為：{', '.join(required)}")

    names: list[str] = []
    for row in rows:
        primary = str(row[0] or "").strip() if len(row) > 0 else ""
        aliases = str(row[1] or "").strip() if len(row) > 1 else ""
        if primary:
            names.append(primary)
        names.extend(_split_aliases(aliases))
    workbook.close()
    return tuple(dict.fromkeys(names))


@lru_cache(maxsize=4)
def load_mrt_location_names(path: str = str(DEFAULT_MRT_CSV)) -> tuple[str, ...]:
    """读取 cp950 CSV 第二列，并补充不含出口编号的站名。"""
    names: list[str] = []
    with open(path, "r", encoding="cp950", newline="") as handle:
        reader = csv.reader(handle)
        next(reader, None)
        for row in reader:
            if len(row) < 2:
                continue
            raw = str(row[1] or "").strip()
            if not raw:
                continue
            names.append(raw)
            station = re.sub(r"(?:出口.*|[A-Za-z]\d+)$", "", raw).strip()
            if station:
                names.append(station)
    return tuple(dict.fromkeys(names))


def contains_known_name(location: str, names: Iterable[str]) -> bool:
    """地点文字是否包含名单名称，忽略空白与常见标点。"""
    target = _normalized_name(location)
    return bool(target) and any(
        normalized and normalized in target
        for normalized in (_normalized_name(name) for name in names)
    )


def validate_landmark(
    location: str,
    *,
    path: str = str(DEFAULT_LANDMARK_XLSX),
) -> bool:
    """地標驗測（2026-09-01 起改走 addrCheck type=Landmark；path 參數保留相容）。"""
    return verify_address_status(location, "landmark")[0] is True


def validate_mrt(
    location: str,
    *,
    path: str = str(DEFAULT_MRT_CSV),
) -> bool:
    """捷運驗測（2026-09-01 起改走 addrCheck type=Landmark；path 參數保留相容）。"""
    return verify_address_status(location, "mrt")[0] is True
