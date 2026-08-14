from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

try:
    from openpyxl import Workbook, load_workbook
except ImportError:  # openpyxl 僅地標名單載入需要；缺少時跳過相關測試
    Workbook = load_workbook = None

from location_validation_119 import (
    AddrCheckResult,
    DEFAULT_LANDMARK_XLSX,
    LOCATION_TYPE_TO_API_TYPE,
    load_landmark_names,
    load_mrt_location_names,
    verify_address,
)


def _mock_response(payload: dict) -> MagicMock:
    response = MagicMock()
    response.read.return_value = json.dumps(payload).encode("utf-8")
    return response


@unittest.skipIf(load_workbook is None, "需要 openpyxl 才能測試地標名單載入")
class LocationNameLoaderTests(unittest.TestCase):
    """捷運／地標名單載入仍供地點類別分類使用（與驗測分開）。"""

    def tearDown(self) -> None:
        load_landmark_names.cache_clear()
        load_mrt_location_names.cache_clear()

    def test_landmark_template_has_expected_columns(self) -> None:
        workbook = load_workbook(DEFAULT_LANDMARK_XLSX, read_only=True)
        headers = [cell.value for cell in next(workbook.active.iter_rows())]
        workbook.close()
        self.assertEqual(headers, ["地標名稱", "別名", "地址"])

    def test_landmark_loader_supports_aliases(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "landmarks.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["地標名稱", "別名", "地址"])
            sheet.append(["板橋大觀市場", "大觀市場、板橋市場", "新北市板橋區"])
            workbook.save(path)

            names = load_landmark_names(str(path))
            self.assertIn("板橋市場", names)
            self.assertIn("大觀市場", names)

    def test_real_mrt_csv_uses_second_column_cp950(self) -> None:
        names = load_mrt_location_names()
        self.assertIn("頂埔站出口1", names)


class VerifyAddressTests(unittest.TestCase):
    def test_valid_house_returns_status_true(self) -> None:
        payload = {
            "status": True,
            "reason": "valid",
            "address": "新北市板橋區中山路二段212巷4之2號",
            "hint": "地址有效：新北市板橋區中山路二段212巷4之2號",
        }
        with patch("location_validation_119.urlopen") as mocked:
            mocked.return_value.__enter__.return_value = _mock_response(payload)
            result = verify_address(
                "板橋中山路二段212巷4-2號5樓", location_type="address"
            )

        self.assertEqual(
            result,
            AddrCheckResult(
                status=True,
                reason="valid",
                address="新北市板橋區中山路二段212巷4之2號",
                hint=payload["hint"],
            ),
        )

    def test_sends_bearer_token_and_mapped_type(self) -> None:
        captured = {}

        def _fake_urlopen(request, timeout=None):
            captured["url"] = request.full_url
            captured["type"] = json.loads(request.data.decode("utf-8")).get("type")
            captured["auth"] = request.headers.get("Authorization")
            ctx = MagicMock()
            ctx.__enter__.return_value = _mock_response(
                {"status": False, "reason": "not_found", "hint": "查無"}
            )
            return ctx

        with patch("location_validation_119.urlopen", side_effect=_fake_urlopen):
            with patch.dict(
                "os.environ",
                {"ADDRCHECK_API_TOKEN": "secret", "ADDRCHECK_API_URL": "http://h:8060"},
                clear=False,
            ):
                verify_address("三芝區北海路二段忠孝街口", location_type="intersection")

        self.assertEqual(captured["url"], "http://h:8060/api/AddrCheck/Verify")
        self.assertEqual(captured["type"], "Crossroad")
        self.assertEqual(captured["auth"], "Bearer secret")

    def test_mrt_maps_to_landmark(self) -> None:
        self.assertEqual(LOCATION_TYPE_TO_API_TYPE["mrt"], "Landmark")

    def test_not_found_returns_status_false_with_hint(self) -> None:
        payload = {"status": False, "reason": "not_found", "hint": "查無「中山亂路」"}
        with patch("location_validation_119.urlopen") as mocked:
            mocked.return_value.__enter__.return_value = _mock_response(payload)
            result = verify_address("中山亂路", location_type="address")
        self.assertIs(result.status, False)
        self.assertEqual(result.reason, "not_found")
        self.assertEqual(result.hint, "查無「中山亂路」")

    def test_network_error_yields_status_none(self) -> None:
        with patch(
            "location_validation_119.urlopen",
            side_effect=TimeoutError("timeout"),
        ):
            result = verify_address("新北市板橋區", location_type="address")
        self.assertIsNone(result.status)
        self.assertIn("timeout", result.error or "")

    def test_empty_address_is_unparsable_without_call(self) -> None:
        with patch("location_validation_119.urlopen") as mocked:
            result = verify_address("   ", location_type="address")
        mocked.assert_not_called()
        self.assertIs(result.status, False)
        self.assertEqual(result.reason, "unparsable")


if __name__ == "__main__":
    unittest.main()
